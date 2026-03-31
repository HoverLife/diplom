import argparse
import csv
import json
import re
import sys
from collections import Counter
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import fitz
import pandas as pd
import pytesseract
from dateutil import parser as dt_parser
from openpyxl.styles import Font, PatternFill
from PIL import Image

try:
    from PyQt5 import QtCore, QtGui, QtWidgets
except Exception:
    QtCore = QtGui = QtWidgets = None

# ----------------------------
# Parsing constants
# ----------------------------
KNOWN_BRANDS = {
    "aoc": "AOC",
    "samsung": "Samsung",
    "lg": "LG",
    "dell": "Dell",
    "hp": "HP",
    "acer": "Acer",
    "philips": "Philips",
    "benq": "BenQ",
    "asus": "ASUS",
    "lenovo": "Lenovo",
    "viewsonic": "ViewSonic",
    "sony": "Sony",
    "iiyama": "iiyama",
}

BASE_COLUMNS = [
    "SourceFile",
    "SourcePage",
    "ProcessedAt",
    "LotPosition",
    "AuctionCity",
    "AuctionEndDateTime",
    "ItemType",
    "Brand",
    "Model",
    "Diagonal",
    "SerialNumber",
    "StartPriceRUB",
    "PickupAddress",
    "CityPickup",
    "VideoInputs",
    "PowerType",
    "PowerIncluded",
    "MatrixIssues",
    "Defect1",
    "Defect2",
    "Photo1",
    "Photo2",
    "Photo3",
    "HasImages",
    "RawText",
    "ExtractedFeatures",
    "Keywords",
    "EstimatedCondition",
    "Notes",
]

KEY_TO_COLUMN = {
    "видеовходы": "VideoInputs",
    "тип питания": "PowerType",
    "сетевой шнур или блок питания в наличии": "PowerIncluded",
    "затемнения и полосы на матрице": "MatrixIssues",
    "диагональ": "Diagonal",
    "дефект1": "Defect1",
    "дефект2": "Defect2",
}

LOT_START_RE = re.compile(r"(?m)^\s*\d+\s*-{2,}\s*")
LOT_FALLBACK_RE = re.compile(r"(?m)^\s*\d+\s+")
DATE_RE = re.compile(r"(\d{1,2}[./]\d{1,2}[./]\d{2,4}(?:\s+\d{1,2}:\d{2})?)")
SERIAL_RE = re.compile(r"(?i)s/n\s*:\s*([A-Za-z0-9\-]+)")
ADDRESS_RE = re.compile(r"(г\.\s*[А-ЯЁA-Z][^\n]+)")
CITY_FROM_ADDRESS_RE = re.compile(r"г\.\s*([А-ЯЁA-Zа-яa-z\-]+)")
PRICE_LINE_RE = re.compile(r"^\s*([\d\s]{2,12})\s*(?:руб|р\.|RUB)?\s*$", flags=re.IGNORECASE)


@dataclass
class ParseResult:
    rows: List[Dict]
    errors: List[Dict]


@dataclass
class ParseConfig:
    save_images: bool = True
    force_ocr: bool = False
    ocr_lang: str = "rus+eng"
    promote_feature_threshold: int = 3


# ----------------------------
# Text normalization & parsing
# ----------------------------
def normalize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("—", "-").replace("–", "-")
    text = re.sub(r"[\t ]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def split_lots(page_text: str) -> List[str]:
    text = "\n" + page_text.strip()
    starts = [m.start() for m in LOT_START_RE.finditer(text)]
    if not starts:
        starts = [m.start() for m in LOT_FALLBACK_RE.finditer(text)]
    if not starts:
        return [page_text]

    blocks: List[str] = []
    for i, start in enumerate(starts):
        end = starts[i + 1] if i + 1 < len(starts) else len(text)
        block = text[start:end].strip()
        if block:
            blocks.append(block)
    return blocks


def normalize_diagonal(value: str) -> str:
    m = re.search(r"(\d{1,2}(?:[\.,]\d{1,2})?)\s*['\"]?", value)
    if not m:
        return value.strip()
    return f"{m.group(1).replace(',', '.')}\""


def parse_int_price(lines: List[str]) -> Optional[int]:
    for line in lines:
        m = PRICE_LINE_RE.match(line)
        if not m:
            continue
        digits = re.sub(r"\s+", "", m.group(1))
        if digits.isdigit():
            return int(digits)

    # fallback: numbers with RUB markers inside lines
    for line in lines:
        m = re.search(r"([\d\s]{2,12})\s*(?:руб|р\.|RUB)", line, flags=re.IGNORECASE)
        if m:
            digits = re.sub(r"\s+", "", m.group(1))
            if digits.isdigit():
                return int(digits)
    return None


def parse_datetime_from_header(text: str) -> Optional[datetime]:
    m = DATE_RE.search(text)
    if not m:
        return None
    try:
        return dt_parser.parse(m.group(1), dayfirst=True)
    except Exception:
        return None


def parse_head(block: str) -> Tuple[str, str, str, str, str]:
    city = ""
    item_type = ""
    brand = ""
    model = ""
    diagonal = ""

    city_m = re.search(r"\(([^)]+)\)", block)
    if city_m:
        city = city_m.group(1).strip()

    head_region = block[city_m.end():] if city_m else block
    sn_cut = re.search(r"(?i)\bs/n\s*:", head_region)
    if sn_cut:
        head_text = head_region[: sn_cut.start()].strip()
    else:
        head_text = re.split(r"-{2,}", head_region, maxsplit=1)[0].strip()

    tokens = head_text.split()
    if tokens:
        item_type = tokens[0]

    for i, token in enumerate(tokens[1:], start=1):
        canon = KNOWN_BRANDS.get(token.lower().strip(".,"))
        if canon:
            brand = canon
            if i + 1 < len(tokens):
                model = tokens[i + 1]
                if i + 2 < len(tokens):
                    diagonal = normalize_diagonal(" ".join(tokens[i + 2 :]))
            break

    if not brand and len(tokens) > 1:
        model = " ".join(tokens[1:])

    if not diagonal:
        dm = re.search(r"(\d{1,2}(?:[\.,]\d{1,2})?)\s*['\"]", head_text)
        if dm:
            diagonal = normalize_diagonal(dm.group(0))

    return city, item_type, brand, model, diagonal


def parse_serial(block: str) -> str:
    m = SERIAL_RE.search(block)
    return m.group(1).strip() if m else ""


def parse_address(block: str) -> str:
    m = ADDRESS_RE.search(block)
    return m.group(1).strip() if m else ""


def parse_city_from_address(address: str) -> str:
    m = CITY_FROM_ADDRESS_RE.search(address)
    return m.group(1).strip() if m else ""


def estimate_condition(row: Dict) -> str:
    txt = " ".join([
        row.get("Defect1", ""),
        row.get("Defect2", ""),
        row.get("MatrixIssues", ""),
        row.get("RawText", ""),
    ]).lower()

    bad_markers = ["не работает", "трещ", "бит", "мертв", "полос", "не включается"]
    fair_markers = ["следы эксплуатации", "царап", "потерт", "потёрт"]
    if any(k in txt for k in bad_markers):
        return "Bad"
    if any(k in txt for k in fair_markers):
        return "Fair"
    return "Good"


def derive_keywords(row: Dict) -> str:
    chunks = [
        row.get("Brand", ""),
        row.get("Model", ""),
        row.get("VideoInputs", ""),
        row.get("Defect1", ""),
        row.get("Defect2", ""),
    ]
    out: List[str] = []
    seen = set()
    for chunk in chunks:
        for token in re.split(r"[,/\s]+", chunk):
            t = token.strip()
            low = t.lower()
            if t and len(t) > 2 and low not in seen:
                out.append(t)
                seen.add(low)
    return ", ".join(out[:12])


def split_attributes(attrs_zone: str) -> List[str]:
    # Split only by ' / ' to avoid breaking values like VGA/DVI.
    parts = re.split(r"\s+/\s+", attrs_zone)
    return [p.strip() for p in parts if p.strip()]


def parse_attributes(attrs_zone: str) -> Tuple[Dict[str, str], Dict[str, str], str]:
    mapped: Dict[str, str] = {}
    extra: Dict[str, str] = {}
    free_defect = ""

    for part in split_attributes(attrs_zone):
        m = re.match(r"([А-Яа-яA-Za-z0-9\-\s]+)\s*-\s*(.+)", part)
        if not m:
            if "дефект" in part.lower() or "следы эксплуатации" in part.lower():
                free_defect = part.splitlines()[0].strip(" .")
            continue

        key_raw = m.group(1).strip()
        key = key_raw.lower()
        value = m.group(2).strip(" .")
        mapped_col = KEY_TO_COLUMN.get(key)
        if mapped_col:
            mapped[mapped_col] = normalize_diagonal(value) if mapped_col == "Diagonal" else value
        else:
            extra[key_raw] = value

    return mapped, extra, free_defect


def parse_block(block: str, source_file: str, source_page: int, auction_dt: Optional[datetime], photos: List[str]) -> Tuple[Dict, List[Dict]]:
    row = {c: "" for c in BASE_COLUMNS}
    row.update(
        {
            "SourceFile": source_file,
            "SourcePage": source_page,
            "ProcessedAt": datetime.now().isoformat(timespec="seconds"),
            "AuctionEndDateTime": auction_dt,
            "RawText": block,
            "HasImages": bool(photos),
        }
    )

    for i, photo in enumerate(photos[:3], start=1):
        row[f"Photo{i}"] = photo

    errors: List[Dict] = []

    lot_m = re.match(r"^\s*(\d+)", block)
    row["LotPosition"] = int(lot_m.group(1)) if lot_m else None

    city, item_type, brand, model, diagonal = parse_head(block)
    row["AuctionCity"] = city
    row["ItemType"] = item_type
    row["Brand"] = brand
    row["Model"] = model
    row["Diagonal"] = diagonal
    row["SerialNumber"] = parse_serial(block)

    row["PickupAddress"] = parse_address(block)
    row["CityPickup"] = parse_city_from_address(row["PickupAddress"])

    markers = re.split(r"-{2,}", block, maxsplit=2)
    attrs_zone = markers[2] if len(markers) > 2 else block
    mapped, extras, free_defect = parse_attributes(attrs_zone)
    for k, v in mapped.items():
        row[k] = v

    lines = [ln.strip() for ln in block.split("\n") if ln.strip()]
    row["StartPriceRUB"] = parse_int_price(lines)
    if row["StartPriceRUB"] is None:
        errors.append(
            {
                "SourceFile": source_file,
                "LotPosition": row["LotPosition"],
                "ErrorMessage": "Price not parsed",
                "RawText": block,
            }
        )

    if not row.get("Defect1"):
        row["Defect1"] = free_defect

    row["ExtractedFeatures"] = json.dumps(extras, ensure_ascii=False)
    row["Keywords"] = derive_keywords(row)
    row["EstimatedCondition"] = estimate_condition(row)
    return row, errors


# ----------------------------
# PDF extraction
# ----------------------------
def extract_images(pdf_path: Path, images_root: Path) -> Dict[int, List[str]]:
    out_dir = images_root / pdf_path.stem
    by_page: Dict[int, List[str]] = {}

    with fitz.open(pdf_path) as doc:
        for page_num, page in enumerate(doc, start=1):
            imgs = page.get_images(full=True)
            rel_paths: List[str] = []
            if imgs:
                out_dir.mkdir(parents=True, exist_ok=True)

            for i, img in enumerate(imgs, start=1):
                xref = img[0]
                data = doc.extract_image(xref)
                ext = data.get("ext", "png")
                fname = f"page{page_num:03d}_{i:02d}.{ext}"
                out_file = out_dir / fname
                out_file.write_bytes(data["image"])
                rel_paths.append(str(out_file.relative_to(images_root.parent)))

            by_page[page_num] = rel_paths

    return by_page


def extract_text_from_pdf(pdf_path: Path, force_ocr: bool, ocr_lang: str) -> Tuple[List[Tuple[int, str]], bool, List[str]]:
    pages: List[Tuple[int, str]] = []
    ocr_used = False
    warnings: List[str] = []

    with fitz.open(pdf_path) as doc:
        for page_num, page in enumerate(doc, start=1):
            text = "" if force_ocr else (page.get_text("text") or "")

            if not text.strip():
                pix = page.get_pixmap(dpi=250)
                mode = "RGBA" if pix.alpha else "RGB"
                image = Image.frombytes(mode, [pix.width, pix.height], pix.samples)
                try:
                    text = pytesseract.image_to_string(image, lang=ocr_lang)
                    ocr_used = True
                except pytesseract.TesseractNotFoundError:
                    warnings.append("Tesseract not found in PATH; OCR skipped.")
                    text = ""
                    ocr_used = True
                except Exception as exc:
                    warnings.append(f"OCR failed on page {page_num}: {exc}")
                    text = ""
                    ocr_used = True

            pages.append((page_num, normalize_text(text)))

    return pages, ocr_used, warnings


def promote_frequent_features(rows: List[Dict], threshold: int) -> List[str]:
    counter: Counter = Counter()
    for row in rows:
        raw = row.get("ExtractedFeatures", "{}")
        try:
            features = json.loads(raw) if raw else {}
        except Exception:
            features = {}
        for key in features.keys():
            counter[key] += 1

    promoted = [k for k, c in counter.items() if c >= threshold]

    for row in rows:
        raw = row.get("ExtractedFeatures", "{}")
        try:
            features = json.loads(raw) if raw else {}
        except Exception:
            features = {}

        for key in promoted:
            col = f"Feature_{sanitize_feature_name(key)}"
            row[col] = features.get(key, "")

    return [f"Feature_{sanitize_feature_name(k)}" for k in promoted]


def sanitize_feature_name(name: str) -> str:
    s = re.sub(r"\s+", "_", name.strip())
    s = re.sub(r"[^A-Za-zА-Яа-я0-9_]+", "", s)
    return s[:50] or "Unknown"


def write_excel(df: pd.DataFrame, output_excel: Path):
    with pd.ExcelWriter(output_excel, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Lots")
        ws = writer.book["Lots"]

        # Header style
        header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
        header_font = Font(bold=True)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font

        # Number/date formats + widths
        col_index = {cell.value: idx for idx, cell in enumerate(ws[1], start=1)}
        price_col = col_index.get("StartPriceRUB")
        dt_col = col_index.get("AuctionEndDateTime")

        for idx, col_name in enumerate(df.columns, start=1):
            letter = ws.cell(row=1, column=idx).column_letter
            max_len = max([len(str(col_name))] + [len(str(v)) for v in df[col_name].fillna("").head(500)]) if not df.empty else len(col_name)
            ws.column_dimensions[letter].width = min(70, max(12, int(max_len * 1.1)))

        if price_col:
            for r in range(2, ws.max_row + 1):
                ws.cell(row=r, column=price_col).number_format = "#,##0"

        if dt_col:
            for r in range(2, ws.max_row + 1):
                ws.cell(row=r, column=dt_col).number_format = "YYYY-MM-DD HH:MM"


def write_errors_csv(errors: List[Dict], output_dir: Path):
    err_file = output_dir / "parsing_errors.csv"
    with open(err_file, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=["SourceFile", "LotPosition", "ErrorMessage", "RawText"])
        writer.writeheader()
        if errors:
            writer.writerows(errors)


def process_pdfs(pdf_files: List[Path], output_excel: Path, config: ParseConfig) -> ParseResult:
    rows: List[Dict] = []
    errors: List[Dict] = []
    images_root = output_excel.parent / "output_images"

    for pdf in pdf_files:
        page_images = extract_images(pdf, images_root) if config.save_images else {}
        pages, _ocr_used, ocr_warnings = extract_text_from_pdf(pdf, config.force_ocr, config.ocr_lang)

        for warn in ocr_warnings:
            errors.append({"SourceFile": pdf.name, "LotPosition": None, "ErrorMessage": warn, "RawText": ""})

        auction_dt = parse_datetime_from_header(pages[0][1]) if pages else None

        for page_num, text in pages:
            if not text.strip():
                errors.append({"SourceFile": pdf.name, "LotPosition": None, "ErrorMessage": "Empty page text", "RawText": ""})
                continue

            for block in split_lots(text):
                if len(block.strip()) < 8:
                    continue
                row, row_errors = parse_block(block, pdf.name, page_num, auction_dt, page_images.get(page_num, []))
                rows.append(row)
                errors.extend(row_errors)

    promoted_cols = promote_frequent_features(rows, config.promote_feature_threshold)

    df = pd.DataFrame(rows)
    all_columns = BASE_COLUMNS + promoted_cols
    for col in all_columns:
        if col not in df.columns:
            df[col] = ""
    if df.empty:
        df = pd.DataFrame(columns=all_columns)
    else:
        df = df[all_columns]

    if "StartPriceRUB" in df.columns:
        df["StartPriceRUB"] = pd.to_numeric(df["StartPriceRUB"], errors="coerce")
    if "AuctionEndDateTime" in df.columns:
        df["AuctionEndDateTime"] = pd.to_datetime(df["AuctionEndDateTime"], errors="coerce")

    write_excel(df, output_excel)
    write_errors_csv(errors, output_excel.parent)

    return ParseResult(rows=rows, errors=errors)


# ----------------------------
# Tests
# ----------------------------
def test_on_sample_text() -> bool:
    sample = (
        "1--- (Москва) Монитор AOC M2060SWD2 19,5' s/n: GLXK5HA053007 --- "
        "Диагональ-19,5' / Затемнения и полосы на матрице-Нет / Видеовходы-VGA/DVI / "
        "Тип питания-Кабель / Сетевой шнур или блок питания в наличии - Да / Следы эксплуатации.\n"
        "700\n"
        "г. Москва, шоссе Энтузиастов д.14\n"
    )
    row, errs = parse_block(sample, "example.pdf", 1, None, [])
    checks = [
        row["LotPosition"] == 1,
        row["AuctionCity"] == "Москва",
        row["Brand"] == "AOC",
        row["Model"] == "M2060SWD2",
        row["Diagonal"] == '19.5"',
        row["SerialNumber"] == "GLXK5HA053007",
        row["StartPriceRUB"] == 700,
        row["VideoInputs"] == "VGA/DVI",
        row["PowerType"] == "Кабель",
        row["PowerIncluded"] == "Да",
        row["Defect1"].startswith("Следы эксплуатации"),
        len(errs) == 0,
    ]
    return all(checks)


def test_ocr_fallback_switch() -> bool:
    class DummyPage:
        alpha = False

        def get_text(self, _):
            return ""

        def get_pixmap(self, dpi=250):
            class DummyPix:
                alpha = False
                width = 2
                height = 2
                samples = b"\x00\x00\x00" * 4

            return DummyPix()

    class DummyDoc:
        def __enter__(self):
            return [DummyPage()]

        def __exit__(self, exc_type, exc, tb):
            return False

    old_open = fitz.open
    old_ocr = pytesseract.image_to_string
    try:
        fitz.open = lambda _path: DummyDoc()
        pytesseract.image_to_string = lambda _img, lang="rus+eng": "ocr text"
        pages, used, warns = extract_text_from_pdf(Path("dummy.pdf"), force_ocr=False, ocr_lang="rus+eng")
        return used is True and pages and pages[0][1] == "ocr text" and not warns
    finally:
        fitz.open = old_open
        pytesseract.image_to_string = old_ocr


def run_sample_tests() -> int:
    tests = [
        ("sample lot parsing", test_on_sample_text),
        ("ocr fallback switch", test_ocr_fallback_switch),
    ]
    passed = 0
    print("Running sample parser tests...")
    for name, fn in tests:
        ok = fn()
        print(f"[{'PASS' if ok else 'FAIL'}] {name}")
        if ok:
            passed += 1
    print(f"Tests passed: {passed}/{len(tests)}")
    return 0 if passed == len(tests) else 1


# ----------------------------
# GUI
# ----------------------------
if QtWidgets is not None:
    class DropArea(QtWidgets.QFrame):
        files_dropped = QtCore.pyqtSignal(list)

        def __init__(self):
            super().__init__()
            self.setAcceptDrops(True)
            self.setFrameShape(QtWidgets.QFrame.StyledPanel)
            self.setStyleSheet("QFrame { border: 2px dashed #6f6f6f; border-radius: 8px; background: #fafafa; }")

            layout = QtWidgets.QVBoxLayout(self)
            layout.setContentsMargins(16, 16, 16, 16)
            label = QtWidgets.QLabel("Перетащите PDF сюда")
            label.setAlignment(QtCore.Qt.AlignCenter)
            font = QtGui.QFont()
            font.setPointSize(10)
            label.setFont(font)
            layout.addWidget(label)

        def dragEnterEvent(self, event):
            if event.mimeData().hasUrls():
                event.acceptProposedAction()

        def dropEvent(self, event):
            files = []
            for url in event.mimeData().urls():
                p = Path(url.toLocalFile())
                if p.suffix.lower() == ".pdf":
                    files.append(str(p))
            if files:
                self.files_dropped.emit(files)


    class Worker(QtCore.QThread):
        finished_ok = QtCore.pyqtSignal(object)
        failed = QtCore.pyqtSignal(str)

        def __init__(self, files: List[Path], output: Path, config: ParseConfig):
            super().__init__()
            self.files = files
            self.output = output
            self.config = config

        def run(self):
            try:
                res = process_pdfs(self.files, self.output, self.config)
                self.finished_ok.emit(res)
            except Exception as exc:
                self.failed.emit(str(exc))


    class MainWindow(QtWidgets.QWidget):
        def __init__(self):
            super().__init__()
            self.setWindowTitle("PDF Auction Parser")
            self.setMinimumSize(580, 360)
            self.resize(640, 400)

            self.files: List[Path] = []
            self.worker: Optional[Worker] = None

            root = QtWidgets.QVBoxLayout(self)
            root.setSpacing(10)

            title = QtWidgets.QLabel("Парсер аукционных PDF")
            tfont = QtGui.QFont()
            tfont.setPointSize(12)
            tfont.setBold(True)
            title.setFont(tfont)

            self.info = QtWidgets.QLabel("Файлы не выбраны")
            self.drop_area = DropArea()
            self.drop_area.files_dropped.connect(self.on_files_dropped)

            self.progress = QtWidgets.QProgressBar()
            self.progress.setRange(0, 100)
            self.progress.setValue(0)

            self.btn_choose = QtWidgets.QPushButton("Выбрать PDF")
            self.btn_start = QtWidgets.QPushButton("Обработать")
            self.btn_choose.clicked.connect(self.choose_files)
            self.btn_start.clicked.connect(self.start_processing)

            self.chk_ocr = QtWidgets.QCheckBox("Принудительный OCR")
            self.chk_images = QtWidgets.QCheckBox("Сохранять изображения")
            self.chk_images.setChecked(True)

            controls = QtWidgets.QHBoxLayout()
            controls.addWidget(self.btn_choose)
            controls.addWidget(self.btn_start)
            controls.addStretch(1)
            controls.addWidget(self.chk_ocr)
            controls.addWidget(self.chk_images)

            root.addWidget(title)
            root.addWidget(self.info)
            root.addWidget(self.drop_area, 1)
            root.addWidget(self.progress)
            root.addLayout(controls)

        def on_files_dropped(self, files: List[str]):
            self.files = [Path(f) for f in files]
            self.info.setText(f"Выбрано PDF: {len(self.files)}")
            self.start_processing(auto_from_drop=True)

        def choose_files(self):
            files, _ = QtWidgets.QFileDialog.getOpenFileNames(self, "Выберите PDF", "", "PDF Files (*.pdf)")
            if files:
                self.files = [Path(f) for f in files]
                self.info.setText(f"Выбрано PDF: {len(self.files)}")

        def start_processing(self, auto_from_drop: bool = False):
            if not self.files:
                QtWidgets.QMessageBox.warning(self, "Нет файлов", "Выберите или перетащите PDF файлы.")
                return

            default_name = f"parsed_lots_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            out, _ = QtWidgets.QFileDialog.getSaveFileName(self, "Сохранить Excel", default_name, "Excel (*.xlsx)")
            if not out:
                if auto_from_drop:
                    self.info.setText("Обработка отменена")
                return

            self.setEnabled(False)
            self.progress.setRange(0, 0)  # busy indicator
            config = ParseConfig(save_images=self.chk_images.isChecked(), force_ocr=self.chk_ocr.isChecked())
            self.worker = Worker(self.files, Path(out), config)
            self.worker.finished_ok.connect(self.on_done)
            self.worker.failed.connect(self.on_error)
            self.worker.start()

        def on_done(self, result: ParseResult):
            self.progress.setRange(0, 100)
            self.progress.setValue(100)
            self.setEnabled(True)
            QtWidgets.QMessageBox.information(
                self,
                "Готово",
                f"Лотов обработано: {len(result.rows)}\nОшибок: {len(result.errors)}",
            )

        def on_error(self, msg: str):
            self.progress.setRange(0, 100)
            self.progress.setValue(0)
            self.setEnabled(True)
            QtWidgets.QMessageBox.critical(self, "Ошибка", msg)


# ----------------------------
# CLI / Main
# ----------------------------
def collect_pdf_inputs(inputs: List[str]) -> List[Path]:
    files: List[Path] = []
    for item in inputs:
        p = Path(item)
        if p.is_dir():
            files.extend(sorted(p.glob("*.pdf")))
        elif p.is_file() and p.suffix.lower() == ".pdf":
            files.append(p)
    return files


def cli_mode(args) -> int:
    pdf_files = collect_pdf_inputs(args.inputs)
    if not pdf_files:
        print("PDF files not found")
        return 1

    cfg = ParseConfig(save_images=args.save_images, force_ocr=args.force_ocr)
    result = process_pdfs(pdf_files, Path(args.output), cfg)
    print(f"Done. Lots: {len(result.rows)}, errors: {len(result.errors)}, output: {args.output}")
    return 0


def main():
    parser = argparse.ArgumentParser(description="Windows-ready PDF auction parser (GUI + CLI)")
    parser.add_argument("inputs", nargs="*", help="PDF files or folders (CLI mode)")
    parser.add_argument("-o", "--output", default="parsed_lots.xlsx", help="Output xlsx (CLI mode)")
    parser.add_argument("--cli", action="store_true", help="Run CLI mode")
    parser.add_argument("--force-ocr", action="store_true", help="Force OCR for all pages")
    parser.add_argument("--no-save-images", action="store_false", dest="save_images", help="Disable image extraction")
    parser.add_argument("--run-tests", action="store_true", help="Run built-in tests")
    args = parser.parse_args()

    if args.run_tests:
        raise SystemExit(run_sample_tests())

    if args.cli or args.inputs:
        raise SystemExit(cli_mode(args))

    if QtWidgets is None:
        print("PyQt5 not available. Install dependencies or run in CLI mode.")
        raise SystemExit(1)

    app = QtWidgets.QApplication(sys.argv)
    window = MainWindow()
    window.show()
    raise SystemExit(app.exec_())


if __name__ == "__main__":
    main()
